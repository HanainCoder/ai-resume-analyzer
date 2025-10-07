import streamlit as st
import pdfplumber
import docx
import os
import spacy
from sentence_transformers import SentenceTransformer, util
import torch
from fpdf import FPDF

# -----------------------------
# LOAD MODELS
# -----------------------------
@st.cache_resource
def load_models():
    nlp = spacy.load("en_core_web_sm")
    device = 'cuda' if torch.cuda.is_available() else 'cpu'
    model = SentenceTransformer('sentence-transformers/all-MiniLM-L6-v2', device=device)
    return nlp, model, device

nlp, model, device = load_models()

# -----------------------------
# PAGE CONFIG
# -----------------------------
st.set_page_config(
    page_title="AI Resume Analyzer",
    layout='wide',
    page_icon="🤖"
)

st.title("🤖 AI Resume Analyzer")
st.markdown("### Upload your Resume and Job Description to get an **AI-powered skill match report**!")

# -----------------------------
# SIDEBAR SETTINGS
# -----------------------------
if "threshold" not in st.session_state:
    st.session_state["threshold"] = 70  # Default threshold

st.sidebar.header("⚙️ Settings")
st.session_state["threshold"] = st.sidebar.slider(
    "Semantic Match Threshold (%)",
    50,
    90,
    st.session_state["threshold"]
)
threshold = st.session_state["threshold"] / 100

# -----------------------------
# FILE INPUTS
# -----------------------------
resume_file = st.file_uploader("📄 Upload your Resume", type=["pdf", "txt", "docx"])
job_description = st.text_area("💼 Paste Job Description here..!", height=200)

# -----------------------------
# HELPER FUNCTIONS
# -----------------------------
def extract_text(file):
    ext = os.path.splitext(file.name)[1].lower()
    text = ""

    try:
        if ext == '.pdf':
            with pdfplumber.open(file) as pdf:
                for page in pdf.pages:
                    text += page.extract_text() or ""
        elif ext == '.docx':
            doc = docx.Document(file)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext == ".txt":
            text = file.read().decode("utf-8")
        else:
            st.error("❌ Unsupported file format!")
    except Exception as e:
        st.error(f"Error reading file: {e}")

    return text.strip()


def extract_skills(text):
    skill_keywords = [
        "python", "java", "c++", "javascript", "sql", "html", "css",
        "react", "node", "django", "flask", "machine learning",
        "deep learning", "nlp", "communication", "leadership",
        "teamwork", "problem solving", "data analysis", "project management",
        "tensorflow", "pytorch", "keras"
    ]
    text = text.lower()
    return list(set(skill for skill in skill_keywords if skill in text))


def semantic_similarity(resume_skills, job_skills, threshold=0.7):
    matched, missing = [], []
    for js in job_skills:
        job_emb = model.encode(js, convert_to_tensor=True)
        found = False
        for rs in resume_skills:
            res_emb = model.encode(rs, convert_to_tensor=True)
            similarity = util.pytorch_cos_sim(job_emb, res_emb).item()
            if similarity >= threshold:
                matched.append(js)
                found = True
                break
        if not found:
            missing.append(js)
    return matched, missing


def create_pdf_report(matched, missing, match_percent):
    """Create a downloadable PDF report"""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 18)
    pdf.cell(200, 10, txt="AI Resume Analyzer Report", ln=True, align="C")
    pdf.ln(10)

    pdf.set_font("Arial", "", 12)
    pdf.cell(200, 10, f"Match Percentage: {match_percent:.2f}%", ln=True)
    pdf.ln(5)

    pdf.set_font("Arial", "B", 14)
    pdf.cell(200, 10, "Matched Skills:", ln=True)
    pdf.set_font("Arial", "", 12)
    pdf.multi_cell(0, 8, ", ".join(matched) if matched else "No matches found.")
    pdf.ln(5)

    pdf.set_font("Arial", "B", 14)
    pdf.cell(200, 10, "Missing Skills:", ln=True)
    pdf.set_font("Arial", "", 12)
    pdf.multi_cell(0, 8, ", ".join(missing) if missing else "No missing skills!")
    pdf.ln(10)

    path = "resume_analysis_report.pdf"
    pdf.output(path)
    return path


# -----------------------------
# MAIN ANALYSIS LOGIC
# -----------------------------
if st.button("🚀 Analyze"):
    if resume_file and job_description:
        resume_text = extract_text(resume_file)
        if not resume_text.strip():
            st.error("⚠ No text found in the uploaded resume.")
        else:
            resume_skills = extract_skills(resume_text)
            job_skills = extract_skills(job_description)
            matched, missing = semantic_similarity(resume_skills, job_skills, threshold)
            match_percent = (len(matched) / len(job_skills)) * 100 if job_skills else 0

            st.markdown("## 🧾 Analysis Summary")
            st.progress(int(match_percent))
            st.markdown(f"### 🎯 Match Score: **{match_percent:.2f}%**")

            col1, col2 = st.columns(2)

            with col1:
                st.markdown("### ✅ Matched Skills")
                if matched:
                    st.markdown(
                        " ".join([
                            f"<span style='color:white;background-color:green;padding:5px 10px;border-radius:10px;margin:3px;display:inline-block'>{s}</span>"
                            for s in matched
                        ]), unsafe_allow_html=True)
                else:
                    st.info("No matched skills found.")

            with col2:
                st.markdown("### ❌ Missing Skills")
                if missing:
                    st.markdown(
                        " ".join([
                            f"<span style='color:white;background-color:#d9534f;padding:5px 10px;border-radius:10px;margin:3px;display:inline-block'>{s}</span>"
                            for s in missing
                        ]), unsafe_allow_html=True)
                else:
                    st.success("No missing skills!")

            st.divider()
            st.subheader("🧩 Extracted Resume Skills")
            st.write(", ".join(resume_skills) if resume_skills else "No skills detected.")

            st.subheader("📋 Job Description Skills")
            st.write(", ".join(job_skills) if job_skills else "No skills detected.")

            st.divider()
            with st.expander("📜 View Extracted Resume Text"):
                st.write(resume_text[:2000] + "..." if len(resume_text) > 2000 else resume_text)

            # Generate and download PDF report
            pdf_path = create_pdf_report(matched, missing, match_percent)
            with open(pdf_path, "rb") as pdf_file:
                st.download_button(
                    label="📥 Download Analysis Report (PDF)",
                    data=pdf_file,
                    file_name="resume_analysis_report.pdf",
                    mime="application/pdf",
                )

            st.success("✅ Analysis complete!")
    else:
        st.warning("⚠ Please upload a resume and paste a job description before analyzing.")
